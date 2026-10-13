"""Approximate rendering of an exported PPTX (background image + text boxes) to PNG, to check text placement."""
import sys, io, os
from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont
E = 12192000 / 1920
HIRA = os.environ.get("LECTURE_SUBTITLE_FONT", "/System/Library/Fonts/Hiragino Sans GB.ttc")
MENLO = os.environ.get("LECTURE_MONO_FONT", "/System/Library/Fonts/Menlo.ttc")
prs = Presentation(sys.argv[1])
for idx, s in enumerate(prs.slides):
    im = None
    for sh in s.shapes:
        if sh.shape_type == 13:
            im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB").resize((1920, 1080))
    d = ImageDraw.Draw(im)
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        x, y = sh.left / E, sh.top / E
        yy = y
        for para in sh.text_frame.paragraphs:
            runs = para.runs
            if not runs:
                continue
            sz = (runs[0].font.size.pt if runs[0].font.size else 12) * 2
            lh = para.line_spacing.pt * 2 if para.line_spacing is not None and hasattr(para.line_spacing, "pt") else sz * 1.2
            xx = x
            for r in runs:
                mono = (r.font.name or "") == "Menlo"
                f = ImageFont.truetype(MENLO if mono else HIRA, int(sz), index=0 if mono else (2 if r.font.bold else 0))
                col = r.font.color.rgb if r.font.color and r.font.color.type is not None else (0, 0, 0)
                d.text((xx, yy + (lh - sz) / 2), r.text, font=f, fill=tuple(col))
                xx += d.textlength(r.text, font=f)
            yy += lh
    im.save(f"out/pptx_sim_{idx}.png")
print("ok")
