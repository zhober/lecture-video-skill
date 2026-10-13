"""Export the lecture as an editable PPTX: one slide per scene (a page's final step, plus a slide for
every step whose content a later step hides via data-until), narration in the speaker notes.

Each slide = the page rendered with its text hidden (background picture) + one editable text box per
rendered line of text. Two things keep the boxes on top of their cards in PowerPoint:
  * fonts: PowerPoint cannot use PingFang SC / Menlo and silently substitutes, so for the export the
    page is laid out with the Office-bundled fonts (assets/pptx_fonts.css aliases PingFang SC ->
    Microsoft YaHei, Menlo -> Consolas) and the boxes ask for exactly those fonts;
  * line breaks: lines are taken from Chrome's layout and written one box per line with wrapping off,
    so PowerPoint never re-wraps; boxes are placed by baseline (BASELINE_K, calibrated against
    PowerPoint's own PDF export with pptx_check.py).

usage: python3 export_pptx_office.py OUT.pptx [pid ...]
writes OUT.pptx, OUT.pptx.boxes.json (every box with its intended baseline, for pptx_check.py) and
out/pptx/<pid>_s<step>_ref.png (the same layout with text visible, the reference for checking).
"""

import json, re, sys
from pathlib import Path
import os
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
EMU_PX = 12192000 / 1920          # slide is 1920 px wide; 1 px = 0.5 pt

# PowerPoint font names and their ascent as a fraction of the font size (hhea ascender, what Chrome uses
# for the top of a text fragment's box)
FONTS = {"sans": ("Microsoft YaHei", 1.0581), "mono": ("Consolas", 0.7427)}
# distance from a PowerPoint text box's top to the first baseline, in units of font size, for a box with
# zero insets and single line spacing; measured from PowerPoint's PDF export (pptx_check.py --calibrate)
BASELINE_K = {"sans": 0.9884, "mono": 0.9206}
# PowerPoint snaps lines differently per point size; residual baseline error (px) per kind and px size,
# measured by pptx_check.py --calibrate and subtracted from each box's top
CALIB = ROOT / "cache" / "pptx_calibration.json"
SIZE_CORR = json.loads(CALIB.read_text()) if CALIB.exists() else {}
FONT_FILES = {("sans", False): "yahei", ("sans", True): "yahei_bold", ("mono", False): "consolas", ("mono", True): "consolas_bold"}
_ADV = {}


def advances(kind, bold):
    """advance width per character (in em) of the Office font, to predict how wide PowerPoint sets a run"""
    key = (kind, bold)
    if key not in _ADV:
        from fontTools.ttLib import TTFont
        f = TTFont(ROOT / "cache" / "fonts_office" / f"{FONT_FILES[key]}.ttf", lazy=True)
        upm, cmap, hmtx = f["head"].unitsPerEm, f.getBestCmap(), f["hmtx"]
        _ADV[key] = lambda ch: hmtx[cmap[ord(ch)]][0] / upm if ord(ch) in cmap else None
    return _ADV[key]


def size_corr(kind, size):
    """px to move a box up (kind) or to narrow each character (kind + '_w') at this font size"""
    table = SIZE_CORR.get(kind, {})
    if not table:
        return 0.0
    k = min(table, key=lambda x: abs(float(x) - size))
    return table[k] if abs(float(k) - size) < 0.75 else 0.0

COLLECT_JS = r"""(ASC) => {
  const hidden = e => !e || e.closest('[data-s]:not(.on)') || getComputedStyle(e).visibility === 'hidden' || getComputedStyle(e).display === 'none';
  const kind = fam => /Menlo|SF Mono|monospace|Consolas/i.test(fam) ? 'mono' : 'sans';
  const rotated = e => { for (let c = e; c && c !== document.body; c = c.parentElement) { const t = getComputedStyle(c).transform; if (t && t !== 'none') { const m = t.match(/matrix\(([^)]+)\)/); if (m) { const v = m[1].split(',').map(Number); if (Math.abs(v[1]) > 1e-3 || Math.abs(v[2]) > 1e-3 || Math.abs(v[0] - v[3]) > 1e-3) return true; } } } return false; };
  // clip rect from ancestors with overflow hidden
  const clipOf = e => { let x0 = -1e9, y0 = -1e9, x1 = 1e9, y1 = 1e9; for (let c = e.parentElement; c && c !== document.body; c = c.parentElement) { const cs = getComputedStyle(c); if (cs.overflow !== 'visible' || cs.overflowX !== 'visible' || cs.overflowY !== 'visible') { const r = c.getBoundingClientRect(); x0 = Math.max(x0, r.left); y0 = Math.max(y0, r.top); x1 = Math.min(x1, r.right); y1 = Math.min(y1, r.bottom); } } return [x0, y0, x1, y1]; };
  // freeze colours of every element that owns text, so hiding one element never hides a child that stays in the picture
  document.querySelectorAll('body *').forEach(e => { if (!e.closest('svg') && [...e.childNodes].some(n => n.nodeType === 3 && n.textContent.trim())) e.style.color = getComputedStyle(e).color; });

  const frags = [], keep = [];
  const tw = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
  const rg = document.createRange();
  let n;
  while ((n = tw.nextNode())) {
    const el = n.parentElement;
    if (!el || el.closest('script,style') || hidden(el)) continue;
    if (!n.textContent.trim()) continue;
    const inSvg = !!el.closest('svg');
    const cs = getComputedStyle(el);
    const fam = cs.fontFamily;
    if (/KuaiLe/i.test(fam)) continue;                          // decorative font: stays in the picture
    let scale = 1;
    if (inSvg) {
      const te = el.closest('text'); if (!te || hidden(te)) continue;
      const m = te.getScreenCTM(); if (!m || Math.abs(m.b) > 1e-3 || Math.abs(m.c) > 1e-3 || Math.abs(m.a - m.d) > 1e-3 * Math.abs(m.a)) continue;   // rotated, skewed or stretched (e.g. sequence-logo letters): stays in the picture
      scale = m.a;
    } else if (rotated(el)) continue;
    const fill = inSvg ? cs.fill : cs.color;
    if (!fill || fill === 'none' || fill === 'transparent' || /rgba\([^)]*,\s*0\)/.test(fill)) continue;
    const size = parseFloat(cs.fontSize) * scale;
    const k = kind(fam);
    const pre = cs.whiteSpace.startsWith('pre');
    const [cx0, cy0, cx1, cy1] = clipOf(el);
    const s = n.textContent;
    let cur = null;
    const flush = () => { if (cur && cur.text.trim()) frags.push(cur); cur = null; };
    for (let i = 0; i < s.length; i++) {
      const cp = s.codePointAt(i); const w = cp > 0xffff ? 2 : 1;
      let ch = s.slice(i, i + w);
      rg.setStart(n, i); rg.setEnd(n, i + w);
      const q = [...rg.getClientRects()].find(r => r.width > 0.01);
      if (w === 2) i++;
      if (!q) { if (ch === '\n' && pre) flush(); continue; }
      const mx = (q.left + q.right) / 2, my = (q.top + q.bottom) / 2;
      if (mx < cx0 || mx > cx1 || my < cy0 || my > cy1) { flush(); continue; }     // clipped away
      if (ch === '\t') { flush(); continue; }                     // tab stop: PowerPoint would draw a narrow space
      if (/[\u2500-\u257F]/.test(ch)) { keep.push([n, i - w + 1, i + 1]); flush(); continue; }   // box drawing: stays in the picture
      if (ch === '\n' || ch === '\r') ch = ' ';
      const base = q.top + ASC[k] * size;
      if (cur && Math.abs(base - cur.base) < 1.5 && Math.abs(q.left - cur.right) < 1.5) { cur.text += ch; cur.right = q.right; }
      else { flush(); cur = { text: ch, left: q.left, right: q.right, base, size, kind: k,
                              color: fill, weight: cs.fontWeight, ls: parseFloat(cs.letterSpacing) || 0, el,
                              italic: cs.fontStyle === 'italic' || cs.fontStyle === 'oblique',
                              under: (cs.textDecorationLine || '').includes('underline') }; }
    }
    flush();
  }
  // merge fragments that continue each other on the same baseline (e.g. differently coloured spans)
  frags.sort((a, b) => a.base - b.base || a.left - b.left);
  const lines = [];
  for (const f of frags) {
    const L = lines.find(l => Math.abs(l.base - f.base) < 1.5 && Math.abs(f.left - l.right) < 1.5);
    if (L) { L.runs.push(f); L.right = f.right; }
    else lines.push({ runs: [f], left: f.left, right: f.right, base: f.base });
  }
  // characters kept in the picture get their own span with a frozen colour, so hiding the parent leaves them visible
  keep.reverse().forEach(([node, a, b]) => { if (node.parentElement.closest('svg')) return; const r = document.createRange(); r.setStart(node, a); r.setEnd(node, b);
    const sp = document.createElement('span'); sp.style.color = getComputedStyle(node.parentElement).color; r.surroundContents(sp); });
  frags.forEach(f => { const te = f.el.closest('svg') ? f.el.closest('text') : f.el; te.dataset.pptxText = '1'; });
  return lines.map(l => ({ left: l.left, right: l.right, base: l.base,
    runs: l.runs.map(f => ({ text: f.text, color: f.color, weight: f.weight, size: f.size, kind: f.kind, ls: f.ls,
                             italic: f.italic, under: f.under, w: f.right - f.left })) }));
}"""

HIDE_JS = r"""() => {
  document.querySelectorAll('[data-pptx-text]').forEach(e => {
    if (e.tagName.toLowerCase() === 'text') { e.style.fill = 'transparent'; e.querySelectorAll('tspan').forEach(t => t.style.fill = 'transparent'); }
    else e.style.color = 'transparent';
  });
}"""


def init_script():
    """globals every page sees before its own scripts run (progress-bar labels from script.json)"""
    steps = json.loads((ROOT / "script.json").read_text(encoding="utf-8")).get("progress_steps", [])
    return f"window.PROGRESS_STEPS = {json.dumps(steps, ensure_ascii=False)};"


def export_pages_dir(only=None):
    """copies of pages/ that also link the export-only font aliases (same folder depth, so ../assets works)"""
    d = ROOT / "pages_pptx"
    d.mkdir(exist_ok=True)
    for f in (ROOT / "pages").glob("*.html"):
        if only and f.stem not in only:
            continue
        html = f.read_text(encoding="utf-8")
        html = html.replace('<link rel="stylesheet" href="../assets/theme.css">',
                            '<link rel="stylesheet" href="../assets/theme.css"><link rel="stylesheet" href="../assets/pptx_fonts.css">', 1)
        (d / f.name).write_text(html, encoding="utf-8")
    return d


def rgb(css):
    m = re.findall(r"[\d.]+", css or "")
    if len(m) >= 3:
        return RGBColor(int(float(m[0])), int(float(m[1])), int(float(m[2])))
    return RGBColor(0x2B, 0x2D, 0x42)


def set_fonts(run, face):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", face)


def add_line(slide, ln):
    big = max(ln["runs"], key=lambda r: r["size"])
    k = big["kind"]
    top = ln["base"] - BASELINE_K[k] * big["size"] - size_corr(k, big["size"])
    h = big["size"] * 1.45
    w = ln["right"] - ln["left"] + big["size"] * 0.6          # slack only; wrapping is off
    tb = slide.shapes.add_textbox(Emu(int(ln["left"] * EMU_PX)), Emu(int(top * EMU_PX)), Emu(int(w * EMU_PX)), Emu(int(h * EMU_PX)))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    para.line_spacing = 1.0
    para.space_before = para.space_after = Pt(0)
    for r in ln["runs"]:
        run = para.add_run()
        run.text = r["text"]
        f = run.font
        f.size = Emu(int(round(r["size"] * 0.5 * 100)) * 127)    # sz is stored in 1/100 pt
        bold = str(r["weight"]) in ("bold", "600", "700", "800", "900")
        f.bold = bold
        f.italic = bool(r.get("italic"))
        f.underline = bool(r.get("under"))
        f.color.rgb = rgb(r["color"])
        face = FONTS[r["kind"]][0]
        f.name = face
        set_fonts(run, face)
        # character spacing so the run is exactly as wide as Chrome set it (Chrome kerns and rounds
        # advances; PowerPoint uses the nominal advances): spc is in 1/100 pt = 1/50 px, per character
        adv = advances(r["kind"], bold)
        widths = [adv(ch) for ch in r["text"]]
        spc_px = r["ls"]
        if r.get("w") and all(x is not None for x in widths):
            spc_px = (r["w"] - sum(widths) * r["size"]) / len(r["text"])
        spc_px -= size_corr(r["kind"] + "_w", r["size"])
        if abs(spc_px) > 0.01:
            run._r.get_or_add_rPr().set("spc", str(int(round(spc_px * 50))))
    return {"left": ln["left"], "right": ln["right"], "base": ln["base"], "size": big["size"], "kind": k,
            "text": "".join(r["text"] for r in ln["runs"])}


def main(out, only=None):
    script = json.loads((ROOT / "script.json").read_text(encoding="utf-8"))
    if only:
        script["pages"] = [p for p in script["pages"] if p["id"] in only]
    pdir = export_pages_dir(only)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    tmp = ROOT / "out" / "pptx"
    tmp.mkdir(parents=True, exist_ok=True)
    asc = {k: v[1] for k, v in FONTS.items()}
    record = []
    with sync_playwright() as p:
        b = p.chromium.launch(channel=os.environ.get("LECTURE_BROWSER_CHANNEL", "chrome") or None, headless=True)
        pg = b.new_page(viewport={"width": 1920, "height": 1080})
        pg.add_init_script(init_script())
        for page in script["pages"]:
            pid = page["id"]
            final = max(bt.get("step", 1) for bt in page["beats"])
            html = (ROOT / "pages" / f"{pid}.html").read_text(encoding="utf-8")
            # one slide per scene: content hidden later by data-until gets its own slide
            scenes = sorted({int(u) for u in re.findall(r'data-until="(\d+)"', html) if int(u) < final} | {final})
            prev = 0
            for last in scenes:
                pg.goto(f"file://{pdir}/{pid}.html?step={last}")
                pg.wait_for_selector("body[data-ready='1']")
                pg.evaluate("document.fonts.ready.then(() => 1)")
                pg.wait_for_timeout(80)
                bad = pg.evaluate("[...document.fonts].filter(f => f.status === 'error').map(f => f.family)")
                if bad:   # e.g. Chrome's font sanitizer rejecting a file: the layout would silently fall back
                    raise SystemExit(f"export font failed to load on {pid}: {bad}")
                pg.screenshot(path=str(tmp / f"{pid}_s{last}_ref.png"))
                lines = pg.evaluate(COLLECT_JS, asc)
                pg.evaluate(HIDE_JS)
                pg.wait_for_timeout(50)
                bg = tmp / f"{pid}_s{last}_bg.png"
                pg.screenshot(path=str(bg))
                s = prs.slides.add_slide(blank)
                s.shapes.add_picture(str(bg), 0, 0, prs.slide_width, prs.slide_height)
                boxes = [add_line(s, ln) for ln in lines]
                notes = "\n".join(bt["text"] for bt in page["beats"] if prev < bt.get("step", 1) <= last)
                s.notes_slide.notes_text_frame.text = notes
                record.append({"slide": len(prs.slides), "pid": pid, "step": last, "boxes": boxes})
                print("slide", len(prs.slides), pid, f"s{last}", len(boxes), "lines", flush=True)
                prev = last
        b.close()
    prs.save(out)
    Path(str(out) + ".boxes.json").write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")
    print("saved", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2:])
