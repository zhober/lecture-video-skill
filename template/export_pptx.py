"""Export the lecture as an editable PPTX: one slide per page (final step),
background = page rendered with text hidden, text re-created as editable text boxes,
narration in speaker notes.

usage: python3 export_pptx.py OUT.pptx
"""

import io, json, sys
from pathlib import Path
import os
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent
EMU_PX = 12192000 / 1920

COLLECT_JS = r"""() => {
  const out = [];
  const hidden = e => e.closest('[data-s]:not(.on)') || getComputedStyle(e).visibility === 'hidden' || getComputedStyle(e).display === 'none';
  // code blocks: one text box per block, lines of coloured runs
  document.querySelectorAll('.code').forEach(cb => {
    if (hidden(cb)) return;
    const cs = getComputedStyle(cb), r = cb.getBoundingClientRect();
    const lines = [[]];
    const tw = document.createTreeWalker(cb, NodeFilter.SHOW_TEXT);
    let t, prevLine = null;
    const lineOf = e => { let c = e; while (c && c !== cb) { const d = getComputedStyle(c).display; if (d !== 'inline' && d !== 'contents') return c; c = c.parentElement; } return cb; };
    while ((t = tw.nextNode())) {
      if (hidden(t.parentElement)) continue;
      const le = lineOf(t.parentElement);
      if (prevLine !== null && le !== prevLine && lines[lines.length - 1].length) lines.push([]);
      prevLine = le;
      const col = getComputedStyle(t.parentElement).color;
      const parts = t.textContent.split('\n');
      parts.forEach((part, i) => { if (i > 0) lines.push([]); if (part) lines[lines.length - 1].push([part, col]); });
    }
    while (lines.length && !lines[0].length) lines.shift();
    while (lines.length && !lines[lines.length - 1].length) lines.pop();
    out.push({ kind: 'code', lines, x: r.left + parseFloat(cs.paddingLeft), y: r.top + parseFloat(cs.paddingTop),
      w: r.width - parseFloat(cs.paddingLeft) - parseFloat(cs.paddingRight), h: r.height, size: parseFloat(cs.fontSize),
      lh: parseFloat(cs.lineHeight) || parseFloat(cs.fontSize) * 1.5, color: cs.color, weight: cs.fontWeight, family: cs.fontFamily });
    cb.dataset.pptxText = '1';
  });
  // HTML text: one box per block container, runs in DOM order
  const isInline = e => { const d = getComputedStyle(e).display; return d === 'inline' || d === 'inline-block' && !e.children.length && e.textContent.length < 40 && false; };
  const container = e => { let c = e; while (c && c !== document.body && getComputedStyle(c).display === 'inline') c = c.parentElement; return c; };
  const groups = new Map();
  const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT | NodeFilter.SHOW_ELEMENT);
  let n;
  while ((n = walker.nextNode())) {
    if (n.nodeType === 1) {
      if (n.tagName === 'BR') { const c = container(n); if (groups.has(c)) groups.get(c).runs.push(['\n', null, null, null, null]); }
      continue;
    }
    const raw = n.textContent; if (!raw.replace(/\s+/g, '')) continue;
    const el = n.parentElement;
    if (!el || el.closest('svg') || el.closest('script') || el.closest('style') || el.closest('.code') || hidden(el)) continue;
    const c = container(el);
    const r = document.createRange(); r.selectNodeContents(n);
    const rects = [...r.getClientRects()].filter(q => q.width > 1 && q.height > 1);
    if (!rects.length) continue;
    const cs = getComputedStyle(el);
    if (!groups.has(c)) groups.set(c, { runs: [], rects: [] });
    const g = groups.get(c);
    if (cs.whiteSpace.startsWith('pre')) {
      raw.replace(/^\n+|\n+$/g, '').split('\n').forEach((ln, i) => { if (i > 0) g.runs.push(['\n', null, null, null, null]); g.runs.push([ln, cs.color, cs.fontWeight, parseFloat(cs.fontSize), cs.fontFamily]); });
      g.pre = true;
    } else g.runs.push([raw.replace(/\s+/g, ' '), cs.color, cs.fontWeight, parseFloat(cs.fontSize), cs.fontFamily]);
    g.rects.push(...rects);
    el.dataset.pptxText = '1';
  }
  groups.forEach((g, c) => {
    let x0 = 1e9, y0 = 1e9, x1 = -1e9, y1 = -1e9;
    g.rects.forEach(q => { x0 = Math.min(x0, q.left); y0 = Math.min(y0, q.top); x1 = Math.max(x1, q.right); y1 = Math.max(y1, q.bottom); });
    const tops = new Set(g.rects.map(q => Math.round(q.top / 6)));
    const ccs = getComputedStyle(c), cr = c.getBoundingClientRect();
    const align = ccs.textAlign === 'center' ? 'center' : (ccs.textAlign === 'right' || ccs.textAlign === 'end' ? 'right' : 'left');
    const multi = tops.size > 1 || g.runs.some(r => r[0] === '\n');
    // trim leading/trailing spaces of the whole run list
    if (g.runs.length && !g.pre) { g.runs[0][0] = g.runs[0][0].replace(/^ /, ''); const L = g.runs[g.runs.length - 1]; L[0] = L[0].replace(/ $/, ''); }
    let bx = x0, bw = x1 - x0;
    if (multi || align !== 'left') { bx = cr.left + parseFloat(ccs.paddingLeft); bw = cr.width - parseFloat(ccs.paddingLeft) - parseFloat(ccs.paddingRight); if (!multi && align === 'left') bx = x0; }
    out.push({ kind: 'block', runs: g.runs, x: bx, y: y0, w: bw, h: y1 - y0, wrap: multi, align,
      lh: parseFloat(ccs.lineHeight) || null, size: g.runs[0] ? g.runs[0][3] : 24 });
  });
  // SVG text
  document.querySelectorAll('svg text').forEach(tn => {
    if (hidden(tn)) return;
    const t = tn.textContent.trim(); if (!t) return;
    const tr = tn.getAttribute('transform') || '';
    const ptr = tn.closest('[transform*="rotate"], [transform*="scale"]');
    if (/rotate|scale/.test(tr) || (ptr && /rotate\((?!0)/.test(ptr.getAttribute('transform')))) return; // keep rotated/scaled text in image
    const r = tn.getBoundingClientRect(); if (r.width < 1) return;
    const cs = getComputedStyle(tn);
    out.push({ kind: 'svg', text: t, x: r.left, y: r.top, w: r.width, h: r.height, size: parseFloat(cs.fontSize),
      color: cs.fill, weight: cs.fontWeight, family: cs.fontFamily, align: 'left', pre: false });
    tn.dataset.pptxText = '1';
  });
  return out;
}"""
HIDE_JS = r"""() => {
  document.querySelectorAll('[data-pptx-text]').forEach(e => {
    if (e.tagName.toLowerCase() === 'text') e.style.fill = 'transparent';
    else { e.style.color = 'transparent'; e.querySelectorAll('*').forEach(c => { if (!c.closest('svg')) c.style.color = 'transparent'; }); }
  });
}"""

def init_script():
    """globals every page sees before its own scripts run (progress-bar labels from script.json)"""
    steps = json.loads((ROOT / "script.json").read_text(encoding="utf-8")).get("progress_steps", [])
    return f"window.PROGRESS_STEPS = {json.dumps(steps, ensure_ascii=False)};"


def set_ea(run, face="PingFang SC"):
    from pptx.oxml.ns import qn
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        pass
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", face)


def rgb(css):
    import re
    m = re.findall(r"[\d.]+", css or "")
    if len(m) >= 3:
        return RGBColor(int(float(m[0])), int(float(m[1])), int(float(m[2])))
    return RGBColor(0x2B, 0x2D, 0x42)


def main(out, only=None):
    script = json.loads((ROOT / "script.json").read_text(encoding="utf-8"))
    if only:
        script["pages"] = [p for p in script["pages"] if p["id"] in only]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    tmp = ROOT / "out" / "pptx"
    tmp.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        b = p.chromium.launch(channel=os.environ.get("LECTURE_BROWSER_CHANNEL", "chrome") or None, headless=True)
        pg = b.new_page(viewport={"width": 1920, "height": 1080})
        pg.add_init_script(init_script())
        for page in script["pages"]:
            pid = page["id"]
            last = max(bt.get("step", 1) for bt in page["beats"])
            pg.goto(f"file://{ROOT}/pages/{pid}.html?step={last}")
            pg.wait_for_selector("body[data-ready='1']")
            pg.evaluate("document.fonts.ready")
            items = pg.evaluate(COLLECT_JS)
            pg.evaluate(HIDE_JS)
            pg.wait_for_timeout(50)
            bg = tmp / f"{pid}_bg.png"
            pg.screenshot(path=str(bg))
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(str(bg), 0, 0, prs.slide_width, prs.slide_height)
            for it in items:
                if it["kind"] == "code":
                    tb = s.shapes.add_textbox(Emu(int(it["x"] * EMU_PX)), Emu(int(it["y"] * EMU_PX)),
                                              Emu(int((it["w"] + 20) * EMU_PX)), Emu(int(it["h"] * EMU_PX)))
                    tf = tb.text_frame
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.word_wrap = False
                    tf.vertical_anchor = MSO_ANCHOR.TOP
                    for li, runs in enumerate(it["lines"]):
                        para = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                        para.line_spacing = Pt(it["lh"] * 0.5)
                        if not runs:
                            runs = [[" ", it["color"]]]
                        for txt_, col in runs:
                            run = para.add_run(); run.text = txt_
                            run.font.size = Pt(it["size"] * 0.5); run.font.name = "Menlo"; run.font.color.rgb = rgb(col); set_ea(run)
                    continue
                if it["kind"] == "block":
                    pad = 3
                    tb = s.shapes.add_textbox(Emu(int((it["x"] - pad) * EMU_PX)), Emu(int((it["y"] - pad) * EMU_PX)),
                                              Emu(int((it["w"] + 2 * pad + (0 if it["wrap"] else it["size"] * 0.8)) * EMU_PX)), Emu(int((it["h"] + 2 * pad) * EMU_PX)))
                    tf = tb.text_frame
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.word_wrap = bool(it["wrap"])
                    tf.vertical_anchor = MSO_ANCHOR.TOP
                    para = tf.paragraphs[0]
                    al = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(it["align"], PP_ALIGN.LEFT)
                    para.alignment = al
                    if it.get("lh"):
                        para.line_spacing = Pt(it["lh"] * 0.5)
                    for txt_, col, wt, sz, fam in it["runs"]:
                        if txt_ == "\n":
                            para = tf.add_paragraph(); para.alignment = al
                            if it.get("lh"):
                                para.line_spacing = Pt(it["lh"] * 0.5)
                            continue
                        run = para.add_run(); run.text = txt_
                        f = run.font
                        f.size = Pt(max(6, sz * 0.5)); f.color.rgb = rgb(col)
                        f.bold = str(wt) in ("bold", "600", "700", "800", "900")
                        f.name = "Menlo" if ("Menlo" in fam or "monospace" in fam) else "PingFang SC"
                        set_ea(run)
                    continue
                pad = 2
                tb = s.shapes.add_textbox(Emu(int((it["x"] - pad) * EMU_PX)), Emu(int((it["y"] - pad) * EMU_PX)),
                                          Emu(int((it["w"] + 2 * pad + it["size"] * 0.6) * EMU_PX)), Emu(int((it["h"] + 2 * pad) * EMU_PX)))
                tf = tb.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = bool(it.get("wrap"))
                tf.vertical_anchor = MSO_ANCHOR.TOP
                lines = it["text"].split("\n") if it["pre"] else [it["text"]]
                for li, line in enumerate(lines):
                    para = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                    run = para.add_run()
                    run.text = line
                    f = run.font
                    f.size = Pt(max(6, it["size"] * 0.5))
                    f.bold = str(it["weight"]) in ("bold", "600", "700", "800", "900")
                    f.color.rgb = rgb(it["color"])
                    mono = "Menlo" in it["family"] or "monospace" in it["family"]
                    f.name = "Menlo" if mono else "PingFang SC"
            notes = "\n".join(bt["text"] for bt in page["beats"])
            s.notes_slide.notes_text_frame.text = notes
            print("slide", pid, len(items), "text boxes", flush=True)
        b.close()
    prs.save(out)
    print("saved", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2:])
